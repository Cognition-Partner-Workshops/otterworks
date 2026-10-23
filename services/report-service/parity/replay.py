"""Record or replay the report-service Java baseline corpus."""
import argparse
import hashlib
import http.client
import io
import json
import re
import sys
import time
import urllib.parse
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader

UTC = timezone.utc  # noqa: UP017
HERE = Path(__file__).resolve().parent
CONTEXTS = ["common", "reports", "outputs", "reconciliation", "archive"]
DISPLAY_STAMP = re.compile(r"^[A-Z][a-z]{2} \d{2}, \d{4} \d{2}:\d{2}$")
FILENAME_STAMP = re.compile(r"_(\d{8}_\d{6})(\.(?:csv|pdf|xlsx))(?=$|[\";])")
ISO_STAMP = re.compile(
    r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:\d{2})$"
)
DISPLAY_TEXT = re.compile(r"(?<![A-Za-z])([A-Z][a-z]{2} \d{2}, \d{4} \d{2}:\d{2})(?!\d)")
NOW_SLOP = timedelta(minutes=10)
CHICAGO = ZoneInfo("America/Chicago")


class ReplayError(Exception):
    """A request or response could not be replayed as specified."""


def input_files(directory):
    paths = [
        directory / "docker-compose.yml",
        directory / "requests.json",
        directory / "stub" / "upstream_stub.py",
    ]
    paths.extend(sorted((directory / "fixtures").rglob("*")))
    return sorted(path for path in paths if path.is_file())


def checksum_files(directory, include_reference=True):
    paths = input_files(directory)
    reference = directory / "java-reference.json"
    if include_reference and reference.is_file():
        paths.append(reference)
    return sorted(paths)


def verify_checksums(directory, include_reference=True):
    sums = directory / "SHA256SUMS"
    if not sums.is_file():
        return ["SHA256SUMS is missing"]
    failures = []
    seen = set()
    for line in sums.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            expected, name = line.split(None, 1)
        except ValueError:
            failures.append(f"malformed checksum line: {line!r}")
            continue
        name = name.strip().lstrip("*")
        if not include_reference and name == "java-reference.json":
            continue
        path = directory / name
        seen.add(name)
        if not path.is_file():
            failures.append(f"{name}: file is missing")
            continue
        actual = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual != expected:
            failures.append(f"{name}: expected {expected}, got {actual}")
    required = {path.relative_to(directory).as_posix() for path in input_files(directory)}
    if include_reference:
        required.add("java-reference.json")
    missing = required - seen
    if missing:
        failures.extend(f"{name}: no checksum entry" for name in sorted(missing))
    if include_reference and not (directory / "java-reference.json").is_file():
        failures.append("java-reference.json is missing")
    return failures


def regenerate_checksums(directory):
    lines = []
    for path in checksum_files(directory):
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        lines.append(f"{digest}  {path.relative_to(directory).as_posix()}")
    (directory / "SHA256SUMS").write_text("\n".join(lines) + "\n", encoding="utf-8")


def media_type_family(media_type):
    media_type = (media_type or "").split(";", 1)[0].strip().lower()
    if media_type == "application/json" or media_type.endswith("+json"):
        return "application/json"
    return media_type


def content_type_parts(value):
    if not value:
        return "", {}
    parts = [part.strip() for part in value.split(";")]
    params = {}
    for part in parts[1:]:
        if "=" not in part:
            continue
        key, param_value = part.split("=", 1)
        params[key.strip().lower()] = param_value.strip().strip("\"'").lower()
    return parts[0].lower(), params


def compare_content_type(java_type, target_type, strict):
    problems = []
    java_media, java_params = content_type_parts(java_type)
    target_media, target_params = content_type_parts(target_type)
    if strict:
        if (java_type or "").strip().lower() != (target_type or "").strip().lower():
            problems.append(f"Content-Type {java_type or None} != {target_type or None}")
        return problems
    if media_type_family(java_media) != media_type_family(target_media):
        problems.append(f"media type {java_media or None} != {target_media or None}")
    if java_params.get("charset") != target_params.get("charset"):
        problems.append(
            f"charset {java_params.get('charset') or None} != {target_params.get('charset') or None}"
        )
    other_java = {key: value for key, value in java_params.items() if key != "charset"}
    other_target = {key: value for key, value in target_params.items() if key != "charset"}
    if other_java != other_target:
        problems.append(f"Content-Type parameters {other_java} != {other_target}")
    return problems


def substitute(value, captures):
    if isinstance(value, dict):
        return {key: substitute(item, captures) for key, item in value.items()}
    if isinstance(value, list):
        return [substitute(item, captures) for item in value]
    if isinstance(value, str):
        for name, captured in captures.items():
            value = value.replace("{" + name + "}", str(captured))
    return value


def json_field(value, field):
    current = value
    for key in field.split("."):
        if not isinstance(current, dict) or key not in current:
            raise ReplayError(f"response does not contain capture field {field!r}")
        current = current[key]
    return current


def capture_fields(case, response_json, captures):
    for name, field in case.get("capture", {}).items():
        captured = json_field(response_json, field)
        if captured is None:
            raise ReplayError(f"capture {name!r} from {field!r} is null")
        captures[name] = captured


def parse_datetime(value):
    if not isinstance(value, str) or not ISO_STAMP.match(value):
        return None
    try:
        if value.endswith("Z"):
            value = value[:-1] + "+00:00"
        return datetime.fromisoformat(value)
    except ValueError:
        return None


def classify_instant(value, now, expected=None):
    parsed = parse_datetime(value)
    if parsed is None:
        return f"<unrecognized:{value}>"
    instant = parsed.astimezone(UTC)
    target = expected or now
    if abs(instant - target) > NOW_SLOP:
        return f"<unrecognized:{value}>"
    zone = "UTC" if parsed.utcoffset() == timedelta(0) else "SERVER-TZ"
    return f"<now:{value_format(value)}:{zone}>"


def value_format(value):
    if value.endswith("Z"):
        return "yyyy-MM-dd'T'HH:mm:ss'Z'"
    if re.search(r"\.\d+[+-]\d{2}:\d{2}$", value):
        return "yyyy-MM-dd'T'HH:mm:ss.SSSXXX"
    if re.search(r"\.\d+$", value):
        return "yyyy-MM-dd'T'HH:mm:ss.SSS"
    if re.search(r"[+-]\d{2}:\d{2}$", value):
        return "yyyy-MM-dd'T'HH:mm:ssXXX"
    return "yyyy-MM-dd'T'HH:mm:ss"


def parse_display(value):
    if not isinstance(value, str) or not DISPLAY_STAMP.match(value):
        return None
    try:
        return datetime.strptime(value, "%b %d, %Y %H:%M")
    except ValueError:
        return None


def classify_display(value, now, expected=None, label="display"):
    local_value = parse_display(value)
    if local_value is None:
        return f"<unrecognized:{value}>"
    target = expected or now
    utc_candidate = local_value.replace(tzinfo=UTC)
    chicago_candidate = local_value.replace(tzinfo=CHICAGO)
    if abs(utc_candidate - target) <= NOW_SLOP:
        zone = "UTC"
    elif abs(chicago_candidate.astimezone(UTC) - target) <= NOW_SLOP:
        zone = "SERVER-TZ"
    else:
        return f"<unrecognized:{value}>"
    return f"<now:{label}:{zone}>"


def classify_filename(value, now):
    try:
        local_value = datetime.strptime(value, "%Y%m%d_%H%M%S")
    except ValueError:
        return f"<unrecognized:{value}>"
    utc_candidate = local_value.replace(tzinfo=UTC)
    chicago_candidate = local_value.replace(tzinfo=CHICAGO)
    if abs(utc_candidate - now) <= NOW_SLOP:
        zone = "UTC"
    elif abs(chicago_candidate.astimezone(UTC) - now) <= NOW_SLOP:
        zone = "SERVER-TZ"
    else:
        return f"<unrecognized:{value}>"
    return f"<now:yyyyMMdd_HHmmss:{zone}>"


def normalize_download_url(value, captures):
    if not isinstance(value, str):
        return value
    for name, captured in sorted(captures.items(), key=lambda item: len(str(item[1])), reverse=True):
        value = re.sub(
            rf"(/api/v1/reports/){re.escape(str(captured))}(/download)",
            rf"\g<1>{{{name}}}\g<2>",
            value,
        )
    return value


def normalize_json(value, case, captures, now):
    volatile = set(case.get("volatile", []))
    defaults = case.get("defaultedDates", False)
    if isinstance(value, dict):
        result = {}
        for key, item in value.items():
            if key == "id" and any(str(item) == str(captured) for captured in captures.values()):
                name = next(
                    name for name, captured in captures.items() if str(item) == str(captured)
                )
                result[key] = f"<id:{name}>"
            elif key == "downloadUrl":
                result[key] = normalize_download_url(item, captures)
            elif key in volatile and isinstance(item, str):
                if defaults and key == "dateFrom":
                    result[key] = classify_instant(
                        item, now, (now - timedelta(days=30)).replace(microsecond=0)
                    )
                elif defaults and key == "dateTo":
                    result[key] = classify_instant(item, now, now)
                elif key in ("createdAt", "completedAt", "generatedAt", "generated_at", "timestamp"):
                    result[key] = classify_instant(item, now)
                else:
                    result[key] = classify_instant(item, now)
            else:
                result[key] = normalize_json(item, case, captures, now)
        return result
    if isinstance(value, list):
        return [normalize_json(item, case, captures, now) for item in value]
    return value


def normalize_period(value, now):
    parts = re.split(r"\s+(?:to|—|–)\s+", value, maxsplit=1)
    if len(parts) != 2:
        return f"<unrecognized:{value}>"
    start = classify_display(parts[0], now, now - timedelta(days=30))
    end = classify_display(parts[1], now, now)
    return start + " to " + end


def csv_row_spans(text):
    rows = []
    fields = []
    field_start = 0
    quoted = False
    index = 0
    while index < len(text):
        char = text[index]
        if char == '"':
            if quoted and index + 1 < len(text) and text[index + 1] == '"':
                index += 2
                continue
            quoted = not quoted
        elif not quoted and char == ",":
            fields.append((field_start, index))
            field_start = index + 1
        elif not quoted and char in "\r\n":
            fields.append((field_start, index))
            rows.append(fields)
            fields = []
            if char == "\r" and index + 1 < len(text) and text[index + 1] == "\n":
                index += 1
            field_start = index + 1
        index += 1
    if fields or field_start < len(text):
        fields.append((field_start, len(text)))
        rows.append(fields)
    return rows


def csv_cell_value(text, span):
    value = text[span[0] : span[1]]
    if value.startswith('"') and value.endswith('"'):
        return value[1:-1].replace('""', '"')
    return value


def normalize_defaulted_window_columns(text, now):
    rows = csv_row_spans(text)
    header_index = None
    columns = {}
    for index, row in enumerate(rows):
        header = [csv_cell_value(text, span).lower() for span in row]
        if "window_from" in header and "window_to" in header:
            header_index = index
            columns = {
                "window_from": header.index("window_from"),
                "window_to": header.index("window_to"),
            }
            break
    if header_index is None:
        return text

    expected = {
        "window_from": (now - timedelta(days=30)).replace(microsecond=0),
        "window_to": now,
    }
    replacements = []
    for row in rows[header_index + 1 :]:
        for name, target in expected.items():
            column = columns[name]
            if column >= len(row):
                continue
            span = row[column]
            value = csv_cell_value(text, span)
            if not value:
                continue
            normalized = classify_instant(value, now, target)
            raw = text[span[0] : span[1]]
            if raw.startswith('"') and raw.endswith('"'):
                normalized = f'"{normalized}"'
            replacements.append((span[0], span[1], normalized))
    for start, end, replacement in reversed(replacements):
        text = text[:start] + replacement + text[end:]
    return text


def normalize_generated_text(text, case, now):
    volatile = set(case.get("volatile", []))
    defaults = case.get("defaultedDates", False)
    if defaults:
        period_pattern = re.compile(
            r"(?im)(#?\s*Period:\s*)([A-Z][a-z]{2} \d{2}, \d{4} \d{2}:\d{2}"
            r"\s+(?:to|—|–)\s+[A-Z][a-z]{2} \d{2}, \d{4} \d{2}:\d{2})"
        )

        def period_replacement(match):
            return match.group(1) + normalize_period(match.group(2), now)

        text = period_pattern.sub(period_replacement, text)
        text = normalize_defaulted_window_columns(text, now)

    if "generatedAt" in volatile or "generated_at" in volatile:
        iso_generated = re.compile(
            r"(?i)(generated(?:\s+at)?\s+)(\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z)"
        )

        def iso_replacement(match):
            return match.group(1) + classify_instant(match.group(2), now)

        text = iso_generated.sub(iso_replacement, text)

    display_generated = re.compile(
        r"(?i)(generated(?:\s+at)?[:\s]+)([A-Z][a-z]{2} \d{2}, \d{4} \d{2}:\d{2})"
    )

    def display_replacement(match):
        return match.group(1) + classify_display(match.group(2), now)

    return display_generated.sub(display_replacement, text)


def normalize_excel(raw, case, now):
    try:
        workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
    except Exception as exc:
        raise ReplayError(f"could not read Excel workbook: {exc}") from exc
    sheets = []
    for worksheet in workbook.worksheets:
        raw_rows = list(worksheet.iter_rows())
        rows = []
        for row in raw_rows:
            values = []
            generated_row = any(str(cell.value).strip().lower() == "generated:" for cell in row)
            period_row = any(str(cell.value).strip().lower() == "period:" for cell in row)
            for index, cell in enumerate(row):
                value = "" if cell.value is None else str(cell.value)
                if generated_row and index > 0 and value:
                    value = classify_display(value, now)
                elif period_row and index > 0 and value and case.get("defaultedDates", False):
                    value = normalize_period(value, now)
                else:
                    value = normalize_generated_text(value, case, now)
                values.append({"value": value, "type": cell.data_type})
            rows.append(values)
        sheets.append(
            {
                "name": worksheet.title,
                "rows": rows,
                "merged_ranges": sorted(str(item) for item in worksheet.merged_cells.ranges),
                "auto_filter_ref": worksheet.auto_filter.ref,
            }
        )
    workbook.close()
    return json.dumps({"sheets": sheets}, ensure_ascii=False, sort_keys=True)


def normalize_pdf(raw, case, now):
    try:
        reader = PdfReader(io.BytesIO(raw))
        pages = []
        for page in reader.pages:
            extracted = page.extract_text() or ""
            extracted = normalize_generated_text(extracted, case, now)
            box = page.mediabox
            pages.append(
                {
                    "mediabox": [
                        str(box.left),
                        str(box.bottom),
                        str(box.right),
                        str(box.top),
                    ],
                    "rotation": int(page.get("/Rotate", 0)),
                    "text_lines": extracted.splitlines(),
                }
            )
    except Exception as exc:
        raise ReplayError(f"could not read PDF: {exc}") from exc
    return json.dumps({"page_count": len(pages), "pages": pages}, ensure_ascii=False, sort_keys=True)


def normalize_body(raw, content_type, disposition, case, captures, now):
    media = (content_type or "").split(";", 1)[0].strip().lower()
    text = raw.decode("utf-8", errors="replace")
    if not raw:
        return {"kind": "empty", "value": ""}
    if "json" in media:
        try:
            parsed = json.loads(text)
        except ValueError:
            return {
                "kind": "text",
                "value": normalize_generated_text(text, case, now),
            }
        normalized = normalize_json(parsed, case, captures, now)
        return {
            "kind": "json",
            "value": json.dumps(normalized, ensure_ascii=False),
        }
    if "spreadsheetml" in media or (disposition and ".xlsx" in disposition.lower()):
        return {"kind": "excel", "value": normalize_excel(raw, case, now)}
    if media == "application/pdf" or (disposition and ".pdf" in disposition.lower()):
        return {"kind": "pdf", "value": normalize_pdf(raw, case, now)}
    return {
        "kind": "text",
        "value": normalize_generated_text(text, case, now),
    }


def normalize_disposition(value, now):
    if not value:
        return value

    def replacement(match):
        return classify_filename(match.group(1), now) + match.group(2)

    return FILENAME_STAMP.sub(replacement, value)


def send(base, method, path, body, content_type, accept, timeout):
    url = urllib.parse.urlsplit(base)
    if url.scheme not in ("http", "https") or not url.hostname:
        raise ReplayError(f"base URL must be http(s)://host[:port], got {base!r}")
    conn_cls = http.client.HTTPSConnection if url.scheme == "https" else http.client.HTTPConnection
    conn = conn_cls(url.hostname, url.port, timeout=timeout)
    headers = {"Accept": accept or "application/json"}
    payload = None
    if body is not None:
        if isinstance(body, (dict, list)):
            payload = json.dumps(body, ensure_ascii=False).encode("utf-8")
        elif isinstance(body, str):
            payload = body.encode("utf-8")
        else:
            payload = str(body).encode("utf-8")
        headers["Content-Type"] = content_type or "application/json"
    started = time.monotonic()
    try:
        conn.request(method, url.path.rstrip("/") + path, body=payload, headers=headers)
        response = conn.getresponse()
        raw = response.read()
        result = {
            "status": response.status,
            "contentType": response.getheader("Content-Type") or "",
            "contentDisposition": response.getheader("Content-Disposition"),
            "raw": raw,
        }
    finally:
        conn.close()
    return result, round((time.monotonic() - started) * 1000)


def wait_for(base, wait_spec, captures, timeout):
    path = substitute(wait_spec["path"], captures)
    field = wait_spec["field"]
    values = wait_spec["in"]
    deadline = time.monotonic() + float(wait_spec.get("timeoutSeconds", 60))
    while True:
        response, _ = send(base, "GET", path, None, None, "application/json", timeout)
        if response["status"] == 200:
            try:
                parsed = json.loads(response["raw"].decode("utf-8"))
                current = json_field(parsed, field)
                if current in values:
                    return
            except (ValueError, ReplayError):
                pass
        if time.monotonic() >= deadline:
            raise ReplayError(
                f"waitFor timed out at GET {path}: expected {field} in {values}"
            )
        time.sleep(0.5)


def error_response(message):
    return {
        "status": 0,
        "contentType": "",
        "contentDisposition": None,
        "raw": message.encode("utf-8"),
    }


def record_response(raw_response, case, captures, now):
    response_json = None
    if raw_response["raw"]:
        media = (raw_response["contentType"] or "").split(";", 1)[0].lower()
        if "json" in media:
            try:
                response_json = json.loads(raw_response["raw"].decode("utf-8"))
            except ValueError:
                response_json = None
    if case.get("capture"):
        if response_json is None:
            raise ReplayError(f"case {case['id']} requested capture from a non-JSON response")
        capture_fields(case, response_json, captures)

    disposition = normalize_disposition(raw_response["contentDisposition"], now)
    return {
        "status": raw_response["status"],
        "contentType": raw_response["contentType"] or None,
        "contentDisposition": disposition,
        "body": normalize_body(
            raw_response["raw"],
            raw_response["contentType"],
            disposition,
            case,
            captures,
            now,
        ),
    }


def compare(java_response, target_response, strict_media_type):
    problems = []
    if java_response is None:
        return ["missing Java capture"]
    if java_response["status"] != target_response["status"]:
        problems.append(
            f"status {java_response['status']} != {target_response['status']}"
        )
    problems.extend(
        compare_content_type(
            java_response.get("contentType"),
            target_response.get("contentType"),
            strict_media_type,
        )
    )
    if java_response.get("contentDisposition") != target_response.get("contentDisposition"):
        problems.append(
            "Content-Disposition "
            f"{java_response.get('contentDisposition')!r} != "
            f"{target_response.get('contentDisposition')!r}"
        )
    if java_response.get("body") != target_response.get("body"):
        problems.append("body differs")
    return problems


def response_text(response):
    if response is None:
        return "(missing Java capture)"
    body = response.get("body", {})
    return json.dumps(body, ensure_ascii=False, sort_keys=True)


def parse_contexts(values):
    if not values:
        return CONTEXTS[:]
    selected = []
    for value in values:
        for item in value.split(","):
            item = item.strip()
            if item == "all":
                selected.extend(CONTEXTS)
            elif item in CONTEXTS:
                selected.append(item)
            else:
                raise SystemExit(
                    f"unknown context {item!r}; use {', '.join(CONTEXTS)} or all"
                )
    return [context for context in CONTEXTS if context in selected]


def run_context(context, cases, references, base, stage, strict, timeout, record, captures):
    results = []
    first_divergence = None
    for original_case in cases:
        case = original_case
        try:
            if case.get("waitFor"):
                wait_for(base, case["waitFor"], captures, timeout)
            request_path = substitute(case["path"], captures)
            request_body = substitute(case.get("body"), captures)
            raw_response, elapsed = send(
                base,
                case["method"],
                request_path,
                request_body,
                case.get("contentType"),
                case.get("accept", "application/json"),
                timeout,
            )
            actual = record_response(raw_response, case, captures, datetime.now(UTC))
            error = None
        except (ReplayError, OSError, http.client.HTTPException, TimeoutError) as exc:
            raw_response = error_response(str(exc))
            actual = record_response(raw_response, {}, captures, datetime.now(UTC))
            elapsed = 0
            error = str(exc)

        if record:
            references[case["id"]] = actual
            problems = []
        else:
            expected = references.get(case["id"])
            problems = compare(expected, actual, strict)
        if error:
            problems.insert(0, f"request failed: {error}")
        result = {
            "id": case["id"],
            "method": case["method"],
            "path": case["path"],
            "match": not problems,
            "problems": problems,
            "elapsedMs": elapsed,
            "java": references.get(case["id"]) if not record else actual,
            "target": actual,
        }
        results.append(result)
        if problems and first_divergence is None:
            first_divergence = result
            if stage == "first":
                break
    total = len(cases)
    identical = sum(1 for result in results if result["match"])
    return {
        "context": context,
        "stage": "record" if record else stage,
        "casesInCorpus": total,
        "casesRun": len(results),
        "identical": identical,
        "different": len(results) - identical,
        "firstDivergence": first_divergence["id"] if first_divergence else None,
        "divergentIds": [result["id"] for result in results if not result["match"]],
        "status": (
            "recorded"
            if record
            else (
                "stopped at " + first_divergence["id"]
                if first_divergence and stage == "first"
                else ("all identical" if not first_divergence else f"{len(results) - identical} different")
            )
        ),
        "cases": results,
    }


def summary_lines(base, stage, evidence, started):
    lines = [
        "# report-service replay against the recorded Java responses",
        "",
        f"Target: `{base}`",
        "",
        "Java: `java-reference.json` (recorded capture)",
        "",
        f"Stage: `{stage}`",
        "",
        f"Started (UTC): {started}",
        "",
        "| Context | Cases in corpus | Cases run | Identical | Different | First divergence | Result |",
        "|---|---:|---:|---:|---:|---|---|",
    ]
    for item in evidence:
        lines.append(
            f"| {item['context']} | {item['casesInCorpus']} | {item['casesRun']} | "
            f"{item['identical']} | {item['different']} | {item['firstDivergence'] or '-'} | "
            f"{item['status']} |"
        )
    total_run = sum(item["casesRun"] for item in evidence)
    total_same = sum(item["identical"] for item in evidence)
    lines.extend(["", f"{total_same}/{total_run} replayed cases identical."])
    for item in evidence:
        for result in item["cases"]:
            if result["match"]:
                continue
            lines.extend(
                [
                    "",
                    f"## {result['id']} ({item['context']})",
                    "",
                    f"`{result['method']} {result['path'][:120]}`: "
                    + "; ".join(result["problems"]),
                    "",
                    "Java (recorded):",
                    "```",
                    response_text(result["java"]),
                    "```",
                    "Target:",
                    "```",
                    response_text(result["target"]),
                    "```",
                ]
            )
    return lines


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", required=True, help="target base URL, such as http://localhost:18091")
    parser.add_argument(
        "--context",
        action="append",
        help="common|reports|outputs|reconciliation|archive|all (repeatable or comma separated)",
    )
    parser.add_argument("--stage", choices=["first", "full"], default="full")
    parser.add_argument("--requests", default=str(HERE / "requests.json"))
    parser.add_argument("--java-capture", default=str(HERE / "java-reference.json"))
    parser.add_argument("--out", default=str(HERE / "out"), help="directory for context JSON and SUMMARY.md")
    parser.add_argument("--timeout", type=float, default=30.0, help="per-request timeout in seconds")
    parser.add_argument(
        "--strict-media-type",
        action="store_true",
        help="compare the complete Content-Type value instead of JSON media-type families",
    )
    parser.add_argument("--record", action="store_true", help="record responses and regenerate java-reference.json")
    parser.add_argument("--skip-checksum", action="store_true", help="skip SHA256SUMS verification")
    args = parser.parse_args()

    requests_path = Path(args.requests).resolve()
    directory = requests_path.parent
    if not args.skip_checksum:
        failures = verify_checksums(directory, include_reference=not args.record)
        if failures and (not args.record or (directory / "SHA256SUMS").exists()):
            print("corpus checksum mismatch, refusing to continue:", *failures, sep="\n  ", file=sys.stderr)
            return 3
        if not args.record:
            print("corpus checksums: OK (SHA256SUMS)")
        elif not failures:
            print("corpus input checksums: OK (SHA256SUMS)")

    try:
        cases = json.loads(requests_path.read_text(encoding="utf-8"))["cases"]
    except (OSError, ValueError, KeyError) as exc:
        print(f"cannot read request corpus: {exc}", file=sys.stderr)
        return 3

    selected = parse_contexts(args.context)
    if args.record:
        references = {}
        capture_path = Path(args.java_capture).resolve()
        if capture_path.is_file():
            try:
                existing = json.loads(capture_path.read_text(encoding="utf-8"))
                references = {item["id"]: item["response"] for item in existing.get("cases", [])}
            except (ValueError, KeyError) as exc:
                print(f"cannot read existing Java capture: {exc}", file=sys.stderr)
                return 3
    else:
        try:
            reference_data = json.loads(Path(args.java_capture).read_text(encoding="utf-8"))
            references = {item["id"]: item["response"] for item in reference_data["cases"]}
        except (OSError, ValueError, KeyError) as exc:
            print(f"cannot read Java capture: {exc}", file=sys.stderr)
            return 3

    started = datetime.now(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    evidence = []
    captures = {}
    for context in selected:
        context_cases = [case for case in cases if case["context"] == context]
        result = run_context(
            context,
            context_cases,
            references,
            args.base,
            args.stage,
            args.strict_media_type,
            args.timeout,
            args.record,
            captures,
        )
        result.update({"base": args.base, "startedUtc": started})
        (out / f"{context}.json").write_text(
            json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
        )
        evidence.append(result)
        print(
            f"[{context}] {result['identical']}/{result['casesRun']} identical of "
            f"{result['casesInCorpus']} in corpus; {result['status']}"
        )
        if args.stage == "first" and result["firstDivergence"]:
            result_case = result["cases"][-1]
            print(
                f"  first divergence {result_case['id']}: "
                f"{'; '.join(result_case['problems'])}"
            )

    if args.record:
        capture_path = Path(args.java_capture).resolve()
        ordered = [
            {"id": case["id"], "context": case["context"], "response": references[case["id"]]}
            for case in cases
            if case["id"] in references
        ]
        capture_path.write_text(
            json.dumps({"recordedUtc": started, "cases": ordered}, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        if capture_path.parent == directory:
            regenerate_checksums(directory)

    lines = summary_lines(args.base, "record" if args.record else args.stage, evidence, started)
    (out / "SUMMARY.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print()
    print("\n".join(lines[8 : 8 + len(evidence) + 3]))
    print(f"\nevidence: {out}/<context>.json, summary: {out / 'SUMMARY.md'}")
    if args.record:
        print(f"recorded Java capture: {Path(args.java_capture).resolve()}")
        return 1 if any(item["different"] for item in evidence) else 0
    return 0 if all(item["different"] == 0 for item in evidence) else 1


if __name__ == "__main__":
    sys.exit(main())
