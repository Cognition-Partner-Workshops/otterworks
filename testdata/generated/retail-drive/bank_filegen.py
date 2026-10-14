"""Deterministic fictional banking documents for the concept-demo drive."""
from __future__ import annotations

import csv
import html
import io
import json
import random
import re
import zipfile
from datetime import datetime, timedelta, timezone

from docx import Document
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from bank_taxonomy import CLIENTS, LOANS, MONTHS, PRODUCTS, REGIONS

MIME = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
    "csv": "text/csv",
    "md": "text/markdown",
    "json": "application/json",
    "png": "image/png",
}

FOOTER = "Bank of America Document Center · Concept demo · Fictional data"
_FIXED_TIME = datetime(2024, 1, 1, tzinfo=timezone.utc)
_FIXED_NAIVE_TIME = datetime(2024, 1, 1)


def _rng(name: str, seed: int) -> random.Random:
    return random.Random(f"{seed}:{name}")


def _pick(name: str, values: list[str], r: random.Random) -> str:
    return next((value for value in values if value in name), r.choice(values))


def _facts(name: str, r: random.Random) -> dict[str, object]:
    loan_match = re.search(r"LN-\d{4}-\d{6}", name)
    return {
        "client": _pick(name, CLIENTS, r),
        "loan": loan_match.group(0) if loan_match else r.choice(LOANS),
        "product": _pick(name, PRODUCTS, r),
        "region": _pick(name, REGIONS, r),
        "principal": round(r.uniform(18_000, 1_250_000), 2),
        "rate": round(r.uniform(3.25, 11.75), 3),
        "ltv": round(r.uniform(42, 88), 1),
        "dti": round(r.uniform(14, 43), 1),
        "risk_score": r.randint(1, 5),
        "liquidity_ratio": round(r.uniform(1.08, 1.72), 2),
    }


def _transactions(r: random.Random, count: int = 12) -> list[dict[str, object]]:
    descriptions = [
        "ACH settlement — fictional",
        "Loan payment — demo record",
        "Treasury service fee",
        "Interest credit",
        "Commercial deposit",
        "Electronic transfer — demo",
    ]
    balance = round(r.uniform(22_000, 240_000), 2)
    rows = []
    for i in range(count):
        amount = round(r.uniform(-18_000, 34_000), 2)
        balance = round(balance + amount, 2)
        rows.append({
            "date": (datetime(2025, 1, 1) + timedelta(days=i * 17)).date().isoformat(),
            "reference": f"TXN-DEMO-{i + 1:03d}",
            "description": r.choice(descriptions),
            "amount_usd": amount,
            "balance_usd": balance,
        })
    return rows


def _normalize_office(data: bytes) -> bytes:
    """Fix ZIP member timestamps so office files have reproducible bytes."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with src, zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as dst:
        for info in src.infolist():
            normalized = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            normalized.compress_type = zipfile.ZIP_DEFLATED
            normalized.external_attr = info.external_attr
            normalized.create_system = info.create_system
            normalized.flag_bits = info.flag_bits
            dst.writestr(normalized, src.read(info.filename), compress_type=zipfile.ZIP_DEFLATED,
                         compresslevel=9)
    return out.getvalue()


def _xlsx(name: str, r: random.Random) -> bytes:
    facts = _facts(name, r)
    wb = Workbook()
    wb.properties.title = name
    wb.properties.creator = "Document Center Demo"
    wb.properties.created = _FIXED_NAIVE_TIME
    wb.properties.modified = _FIXED_NAIVE_TIME
    ws = wb.active
    ws.title = "Credit and Liquidity"
    ws.append(["Metric", "Value", "Units / Notes"])
    for label, value, note in [
        ("Fictional client", facts["client"], "Synthetic entity"),
        ("Loan ID", facts["loan"], "Fictional tracking key"),
        ("Product", facts["product"], "Concept-demo product"),
        ("Principal", facts["principal"], "USD"),
        ("Annual rate", facts["rate"], "Percent"),
        ("Loan-to-value", facts["ltv"], "Percent"),
        ("Debt-to-income", facts["dti"], "Percent"),
        ("Credit risk band", facts["risk_score"], "1 = low, 5 = elevated"),
        ("Liquidity coverage ratio", facts["liquidity_ratio"], "Illustrative ratio"),
    ]:
        ws.append([label, value, note])
    activity = wb.create_sheet("Statement Activity")
    activity.append(["Date", "Reference", "Description", "Amount USD", "Balance USD"])
    for tx in _transactions(r, 10):
        activity.append(list(tx.values()))
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 29
        sheet.column_dimensions["B"].width = 34
        sheet.column_dimensions["C"].width = 33
    buf = io.BytesIO()
    wb.save(buf)
    return _normalize_office(buf.getvalue())


def _docx(name: str, r: random.Random) -> bytes:
    facts = _facts(name, r)
    doc = Document()
    doc.core_properties.title = name
    doc.core_properties.author = "Document Center Demo"
    doc.core_properties.last_modified_by = ""
    doc.core_properties.created = _FIXED_TIME
    doc.core_properties.modified = _FIXED_TIME
    doc.add_heading(name, level=0)
    doc.add_paragraph(
        "Fictional concept-demo document for Bank of America Document Center. "
        "All parties and figures in this file are synthetic."
    )
    doc.add_heading("Review Summary", level=1)
    doc.add_paragraph(
        f"{facts['client']} is associated with {facts['loan']} for the "
        f"{facts['product']} product. The illustrative principal is "
        f"${facts['principal']:,.2f} at {facts['rate']:.3f}%."
    )
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Shading Accent 1"
    for label, value in [
        ("Region", facts["region"]),
        ("Loan-to-value", f"{facts['ltv']}%"),
        ("Debt-to-income", f"{facts['dti']}%"),
        ("Credit risk band", facts["risk_score"]),
        ("Liquidity coverage ratio", facts["liquidity_ratio"]),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
    doc.add_heading("Controls and Next Steps", level=1)
    for item in [
        "Confirm the synthetic application packet and supporting checklist are complete.",
        "Record the credit rationale and route exceptions through the documented review path.",
        "Reconcile the illustrative figures before any training or presentation use.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run(FOOTER)
    buf = io.BytesIO()
    doc.save(buf)
    return _normalize_office(buf.getvalue())


def _pptx(name: str, r: random.Random) -> bytes:
    facts = _facts(name, r)
    prs = Presentation()
    prs.core_properties.title = name
    prs.core_properties.author = "Document Center Demo"
    prs.core_properties.last_modified_by = ""
    prs.core_properties.created = _FIXED_TIME
    prs.core_properties.modified = _FIXED_TIME
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = name
    title.placeholders[1].text = "Bank of America Document Center · Concept demo"
    summary = prs.slides.add_slide(prs.slide_layouts[1])
    summary.shapes.title.text = "Illustrative Portfolio Snapshot"
    frame = summary.placeholders[1].text_frame
    frame.text = f"Fictional client: {facts['client']}"
    for line in [
        f"Product: {facts['product']} · Region: {facts['region']}",
        f"Principal: ${facts['principal']:,.2f} · Rate: {facts['rate']:.3f}%",
        f"LTV: {facts['ltv']}% · DTI: {facts['dti']}%",
    ]:
        frame.add_paragraph().text = line
    controls = prs.slides.add_slide(prs.slide_layouts[1])
    controls.shapes.title.text = "Review Controls"
    notes = controls.placeholders[1].text_frame
    notes.text = "All records are synthetic and intended for a concept demo."
    for line in [
        "Validate the source packet and ownership checklist.",
        "Document approval rationale and exception handling.",
        "Review quarterly risk and liquidity indicators.",
    ]:
        notes.add_paragraph().text = line
    buf = io.BytesIO()
    prs.save(buf)
    return _normalize_office(buf.getvalue())


def _footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#525f70"))
    canvas.drawCentredString(LETTER[0] / 2, 0.38 * inch, FOOTER)
    canvas.restoreState()


def _invariant_canvas(*args, **kwargs):
    kwargs["invariant"] = 1
    return pdfcanvas.Canvas(*args, **kwargs)


def _pdf(name: str, r: random.Random) -> bytes:
    facts = _facts(name, r)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CenterNote", parent=styles["Normal"], alignment=TA_CENTER))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, title=name,
                            rightMargin=0.65 * inch, leftMargin=0.65 * inch,
                            topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    rows = [
        ["Review field", "Illustrative value"],
        ["Fictional client", str(facts["client"])],
        ["Loan ID", str(facts["loan"])],
        ["Product / region", f"{facts['product']} / {facts['region']}"],
        ["Principal / rate", f"${facts['principal']:,.2f} / {facts['rate']:.3f}%"],
        ["LTV / DTI", f"{facts['ltv']}% / {facts['dti']}%"],
        ["Credit band / LCR", f"{facts['risk_score']} / {facts['liquidity_ratio']}"],
    ]
    table = Table(rows, colWidths=[1.65 * inch, 4.8 * inch], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#012169")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#aeb8c6")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.white, colors.HexColor("#f1f4f8")]),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story = [
        Paragraph(html.escape(name), styles["Title"]),
        Paragraph("Bank of America Document Center · Concept demo · Fictional data",
                  styles["CenterNote"]),
        Spacer(1, 0.18 * inch),
        Paragraph("Credit and account review", styles["Heading2"]),
        Paragraph(
            f"This synthetic review covers {html.escape(str(facts['client']))} and "
            f"{html.escape(str(facts['loan']))}. Values are illustrative and contain "
            "no real customer or account data.",
            styles["BodyText"],
        ),
        Spacer(1, 0.12 * inch),
        table,
        Spacer(1, 0.18 * inch),
        Paragraph("Control checklist", styles["Heading2"]),
        Paragraph(
            "• Confirm KYC/CDD evidence and beneficial-ownership review.<br/>"
            "• Record credit rationale, LTV and DTI checks.<br/>"
            "• Escalate exceptions through the approved risk workflow.",
            styles["BodyText"],
        ),
    ]
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer,
              canvasmaker=_invariant_canvas)
    return buf.getvalue()


def _csv(name: str, r: random.Random) -> bytes:
    facts = _facts(name, r)
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(["date", "transaction_ref", "fictional_client", "product",
                     "description", "amount_usd", "balance_usd"])
    for tx in _transactions(r):
        writer.writerow([
            tx["date"], tx["reference"], facts["client"], facts["product"],
            tx["description"], f"{tx['amount_usd']:.2f}", f"{tx['balance_usd']:.2f}",
        ])
    return buf.getvalue().encode("utf-8")


def _md(name: str, r: random.Random) -> bytes:
    facts = _facts(name, r)
    lines = [
        f"# {name}",
        "",
        "> Bank of America Document Center concept demo — all figures and entities are fictional.",
        "",
        "## Review summary",
        f"- Fictional client: {facts['client']}",
        f"- Loan reference: {facts['loan']}",
        f"- Product: {facts['product']} ({facts['region']})",
        f"- Illustrative principal: ${facts['principal']:,.2f}",
        f"- Annual rate: {facts['rate']:.3f}%",
        "",
        "## Credit and liquidity controls",
        f"- Loan-to-value: {facts['ltv']}%",
        f"- Debt-to-income: {facts['dti']}%",
        f"- Credit risk band: {facts['risk_score']} (illustrative 1–5 scale)",
        f"- Liquidity coverage ratio: {facts['liquidity_ratio']}",
        "- Verify source records, document exceptions, and retain the review evidence.",
        "",
        "## Synthetic activity sample",
        "| Date | Reference | Description | Amount USD |",
        "|---|---|---|---:|",
    ]
    for tx in _transactions(r, 4):
        lines.append(
            f"| {tx['date']} | {tx['reference']} | {tx['description']} | "
            f"{tx['amount_usd']:.2f} |"
        )
    return ("\n".join(lines) + "\n").encode("utf-8")


def _json(name: str, r: random.Random) -> bytes:
    facts = _facts(name, r)
    obj = {
        "title": name,
        "organization": "Bank of America Document Center",
        "classification": "concept-demo fictional data",
        "client": facts["client"],
        "loan_id": facts["loan"],
        "product": facts["product"],
        "region": facts["region"],
        "principal_usd": facts["principal"],
        "annual_rate_pct": facts["rate"],
        "ltv_pct": facts["ltv"],
        "dti_pct": facts["dti"],
        "credit_risk_band": facts["risk_score"],
        "liquidity_coverage_ratio": facts["liquidity_ratio"],
        "transactions": _transactions(r, 8),
    }
    return (json.dumps(obj, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def _png(name: str, r: random.Random) -> bytes:
    facts = _facts(name, r)
    values = [round(r.uniform(0.75, 1.35) * facts["liquidity_ratio"], 3) for _ in range(8)]
    figure = Figure(figsize=(7, 4), dpi=100)
    FigureCanvasAgg(figure)
    ax = figure.add_subplot(111)
    ax.plot(["Q1", "Q2", "Q3", "Q4", "Q1", "Q2", "Q3", "Q4"],
            values, color="#012169", marker="o", linewidth=2)
    ax.axhline(1.0, color="#dc1431", linestyle="--", linewidth=1, label="Illustrative threshold")
    ax.set_title(name, color="#012169")
    ax.set_ylabel("Illustrative ratio / index")
    ax.set_xlabel("Quarter")
    ax.grid(alpha=0.25)
    ax.legend(fontsize=8)
    figure.tight_layout()
    buf = io.BytesIO()
    figure.savefig(buf, format="png", dpi=100, metadata={"Software": "Document Center Demo"})
    figure.clear()
    return buf.getvalue()


_BUILDERS = {
    "xlsx": _xlsx,
    "docx": _docx,
    "pptx": _pptx,
    "pdf": _pdf,
    "csv": _csv,
    "md": _md,
    "json": _json,
    "png": _png,
}


def build(ext: str, name: str, seed: int, kind: str | None = None) -> tuple[bytes, str]:
    """Return deterministic ``(bytes, mime_type)`` for a supported file type."""
    ext = ext.lower().lstrip(".")
    if ext not in _BUILDERS:
        raise ValueError(f"unsupported bank document type: {ext}")
    if kind not in (None, "chart"):
        raise ValueError(f"unsupported bank document kind: {kind}")
    if kind == "chart" and ext != "png":
        raise ValueError("chart kind is supported only for png files")
    return _BUILDERS[ext](name, _rng(name, seed)), MIME[ext]
